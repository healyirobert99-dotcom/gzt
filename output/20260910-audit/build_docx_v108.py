"""生成 A/H 投研交易工作台 v1.0.8 交付审计文档（docx）。

v1.0.8 是「导入完整性封板修复」版本：不新增业务模块，只修复 v1.0.7 独立代码审计
发现的导入完整性与交付一致性问题。本脚本全新编写，不复用任何历史 docx 字节；
所有测试数字均由本轮实际运行得到（见 §9），不使用任何未验证的声称。
"""

import os
import sys
import datetime
import hashlib

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ================ 基础排版 helper ================
def add_h(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(16)
    elif level == 2:
        run.font.size = Pt(13)
    else:
        run.font.size = Pt(11)


def add_p(doc, text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(11)
    if bold:
        r.bold = True


def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(text)
    r.font.size = Pt(11)


def add_code(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(9)
    r.font.name = 'Consolas'


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
    for ri, row in enumerate(rows, 1):
        for ci, val in enumerate(row):
            cell = t.rows[ri].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)


# ================ 路径 ================
OUT_DIR = r'D:\个股工作台\output\20260910-audit\stage3'
os.makedirs(OUT_DIR, exist_ok=True)
OUT_PATH = os.path.join(OUT_DIR, 'A-H投研交易工作台交付审计文档-v1.0.8.docx')

# ================ 文档主体 ================
doc = Document()
doc.styles['Normal'].font.name = 'Microsoft YaHei'
doc.styles['Normal'].font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('A/H 投研交易工作台')
r.bold = True
r.font.size = Pt(24)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('交付审计文档（v1.0.8 · 导入完整性封板修复）')
r.bold = True
r.font.size = Pt(18)

doc.add_paragraph()
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run(
    '版本：v1.0.8\n'
    '生成时间：' + datetime.datetime.now().strftime('%Y-%m-%d %H:%M') + '\n'
    '基线版本：v1.0.7 → v1.0.8\n'
    '本轮性质：仅修复 v1.0.7 独立代码审计发现的导入完整性与交付一致性问题\n'
    '边界：不新增业务模块、不新增数据库表/列、不改动写入函数 SQL、不改变研究口径\n'
    '例外说明：本轮明确移除「载入示例」按钮，并新增导入 JSON 固定 Schema 校验'
).font.size = Pt(11)

doc.add_page_break()

# ================ §1 本轮定位 ================
add_h(doc, '1. 本轮定位与修复清单', 1)
add_p(doc, 'v1.0.7 的「导入与更新」交互层通过了功能测试，但被独立代码审计指出'
            '导入完整性证明不足与交付包自相矛盾。v1.0.8 只做修复，不扩展功能。')
add_table(doc, ['#', '审计问题', 'v1.0.8 处理'], [
    ['1', 'base64(JSON) token 不能证明 preview 实际发生过',
     '改为服务端 secrets.token_urlsafe 随机 token + 服务端 cache + 5 分钟 TTL + 一次性失效'],
    ['2', '状态确认来自导入 JSON',
     'status_change_confirmed 一律忽略；preview 初次 change_confirmed 恒为 False；确认只来自 commit 请求的 confirmed_status_changes'],
    ['3', 'commit 前未校验 preview snapshot 是否漂移',
     'commit 前重读数据库逐项对照 snapshot；漂移 → 拒绝并要求重新解析预览'],
    ['4', '同一批可重复出现同一证券',
     '_parse_import_full 在 preview 前检查 securities[] 内 (exchange, code) 唯一，重复整体拒绝'],
    ['5', 'core_validations / wall_conditions 结构未严格校验',
     '按 ah-workbench-import v1.0 固定 Schema 严格校验对象结构，不猜测、不自动转换'],
    ['6', '名称不一致只在返回 JSON 中暗藏 name_current',
     'preview 数据显式返回 name_mismatch + name_mismatch_notice，前端页面明显展示；本轮不自动改 name'],
    ['7', '「载入示例」可一键载入真实候选标的的虚构数据',
     '直接删除「载入示例」按钮；样例 JSON 仅保留明显虚构的演示标的并标注"请勿写入正式库"'],
    ['8', '缺少导入完整性回归测试',
     '新增 tests/test_v108.py（本轮 129 断言）+ HTTP 端到端冒烟脚本，覆盖审计要求的 14 类场景'],
    ['9', '交付包无法复现文档声称的断言数、Manifest 与 ZIP 不一致',
     '本版重新实跑全套测试、重建 ZIP 与 Manifest、重算 SHA-256、README 修正默认地址'],
])

doc.add_page_break()

# ================ §2 两段式导入 ================
add_h(doc, '2. 真正落实「两段式导入」（审计问题 1）', 1)
add_p(doc, 'v1.0.7 的 token 是 base64(JSON)。它只是把粘贴内容编了个码，'
            '既不证明 preview 发生过，也无法阻止客户端自行构造。v1.0.8 改为服务端持有状态：')
add_table(doc, ['阶段', '服务端行为'], [
    ['preview', '解析 + 对比生成 diff（不写库）；生成 secrets.token_urlsafe(32) 随机 token；'
                '在服务端 cache 保存 token → {kind, validated, snapshot, created_at, expires_at}；'
                'snapshot 记录涉及证券的 security_id / status / research 最新 version / '
                'trade_plan 最新 version / execution_latest.id'],
    ['commit', 'token 必须命中服务端 cache；不存在 / 已过期 / 已使用 / kind 不匹配 → 拒绝；'
               'commit 成功后 token 立即删除（一次性）'],
])
add_p(doc, '关键性质：', bold=True)
add_bullet(doc, 'token 由服务端随机生成，客户端拿不到也无法伪造；手工构造的 base64(JSON) 必然不在 cache 中。')
add_bullet(doc, 'TTL = 300 秒（与既有文档声明的 5 分钟一致）。')
add_bullet(doc, '代码与文档中不再把 base64(JSON) 称为"预览 token"。')

add_p(doc, '')
add_code(doc,
         "POST /api/import/preview  → {token, token_ttl_seconds: 300, securities: [...], warnings: [...]}\n"
         "POST /api/import/commit   ← {token, confirmed_status_changes: [...]}")

doc.add_page_break()

# ================ §3 状态确认 ================
add_h(doc, '3. 状态确认不得来自导入 JSON（审计问题 2）', 1)
add_bullet(doc, '导入 JSON 中的 status_change_confirmed 被完全忽略（软删除），不写入 validated 结构。')
add_bullet(doc, 'preview 初次返回的所有 status_change 中 change_confirmed 恒为 false。')
add_bullet(doc, '确认只来自 commit 请求体单独提交的 confirmed_status_changes，'
                '支持 [0, 2] 下标形式，或 [{"exchange":"SZ","code":"000001"}] 身份形式。')
add_bullet(doc, '未勾选即 commit → 400 拒绝，且状态不变；decision_ledger 只在确认后写入。')
add_p(doc, '结论：导入 JSON 本身在任何情况下都不能声明"用户已经确认"。')

# ================ §4 漂移检测 ================
add_h(doc, '4. 提交前验证 preview snapshot 未漂移（审计问题 3）', 1)
add_p(doc, 'commit 时重新读库，对 preview 涉及的每只证券逐项比对。命中任一即停止本次 commit：')
add_table(doc, ['漂移类型', '判定'], [
    ['status 改变', 'preview 时 status ≠ 当前 status'],
    ['research 最新 version 改变', 'preview 时 research version ≠ 当前 version'],
    ['trade_plan 最新 version 改变', 'preview 时 trade_plan version ≠ 当前 version'],
    ['原不存在的证券已经出现', 'preview 时 exists=false，现在 exists=true'],
    ['（对称保护）预览时存在、现在消失', 'preview 时 exists=true，现在 exists=false'],
])
add_p(doc, '拒绝时的统一提示：', bold=True)
add_code(doc, '工作台数据自预览后已发生变化，请重新解析预览。\n'
              '<证券标识>：status 由 \'等价格\' 变为 \'可交易\'')
add_p(doc, '实现约束：不得根据新的数据库状态直接继续生成下一版本；'
            '漂移被拒绝时 token 不被消费（用户可重新预览后继续）。')

doc.add_page_break()

# ================ §5 固定 Schema ================
add_h(doc, '5. 研究数组结构固定 Schema（审计问题 5）', 1)
add_p(doc, '自 v1.0.8 起，ah-workbench-import v1.0 的以下两项结构为正式固定 Schema：')
add_code(doc,
         'core_validations[] 每一项必须是对象：\n'
         '  {\n'
         '    "content": "非空文本",\n'
         '    "status": "跟踪中 | 已验证 | 已恶化"\n'
         '  }\n\n'
         'wall_conditions[] 每一项必须是对象：\n'
         '  {\n'
         '    "content": "非空文本",\n'
         '    "triggered": true | false\n'
         '  }')
add_p(doc, '明确拒绝（preview 阶段即 400，不触库）：', bold=True)
add_bullet(doc, '字符串 / 数字 / null / 数组元素（例如 ["xxx", "yyy"]）')
add_bullet(doc, '缺少 content，或 content 为空字符串 / 纯空白')
add_bullet(doc, 'core_validations.status 不在三选一之内')
add_bullet(doc, 'wall_conditions.triggered 非严格布尔（字符串 "false"、数字 0/1 均拒绝）')
add_p(doc, '不猜测、不自动转换错误结构。')

# ================ §6 名称不一致 ================
add_h(doc, '6. 名称不一致必须在预览显式提示（审计问题 6）', 1)
add_bullet(doc, '证券唯一身份只认 exchange + code，绝不用 name 做匹配。')
add_bullet(doc, '导入 name ≠ 数据库 name 时，preview 返回 name_mismatch=true 与可直接渲染的提示文案。')
add_bullet(doc, '前端页面明显展示该提示，而不是只在返回 JSON 中携带 name_current。')
add_bullet(doc, '本轮不自动修改 name。')
add_p(doc, '提示文案：', bold=True)
add_code(doc, '名称不一致：\n工作台：XXX\n导入块：YYY\n证券仍按 exchange+code 识别，请人工核对。')

doc.add_page_break()

# ================ §7 移除危险示例 ================
add_h(doc, '7. 移除危险的真实证券示例（审计问题 7）', 1)
add_bullet(doc, '生产 UI 的「载入示例」按钮已删除，不再存在一键载入虚构研究/交易计划的入口。')
add_bullet(doc, '粘贴页仅保留格式骨架占位提示（placeholder），不提供任何可直接提交的数据。')
add_bullet(doc, 'samples/json/ 下的样例为明显虚构的演示标的（HK.90001 / SH.900001，名称含"（虚构）"），'
                '并带 _notice："仅用于格式演示，请勿写入正式库"。')
add_bullet(doc, '样例中不含任何真实候选标的的虚构首仓区 / 目标仓位等数据。')

# ================ §8 测试结果 ================
add_h(doc, '8. 测试结果（本轮实际运行）', 1)
add_p(doc, '以下数字为本轮在交付源码上实际运行得到，可原样复现；不沿用历史文档的 279 断言声称。')
add_table(doc, ['测试套件', '断言数', '结果', '是否依赖外网'], [
    ['tests/test_v102.py', '38', 'PASS', '否（集成测试单独执行）'],
    ['tests/test_v103.py', '60', 'PASS', '否'],
    ['tests/test_v105.py', '25', 'PASS', '否'],
    ['tests/test_v106.py', '28', 'PASS', '否'],
    ['tests/test_v107.py', '127', 'PASS', '否'],
    ['tests/test_v108.py', '129', 'PASS', '否'],
    ['小计（离线套件）', '407', 'PASS', '—'],
    ['tests/test_integration_quote.py', '14', 'PASS', '是（真实行情接口）'],
    ['output/20260910-audit/smoke_v108_http.py', '14', 'PASS', '否（本机临时端口）'],
])

add_p(doc, '§8.1 test_v108.py 覆盖的 14 类审计要求场景', bold=True)
add_table(doc, ['#', '场景', '用例'], [
    ['1', '不调用 preview 直接 commit → 拒绝', '§A#1'],
    ['2', '手工构造 base64 token → 拒绝', '§A#2'],
    ['3', 'token 成功 commit 后再次使用 → 拒绝', '§A#3'],
    ['4', 'token 超时 → 拒绝', '§A#4'],
    ['5', 'pasted JSON 自带 status_change_confirmed=true → 初次 preview 仍必须未确认', '§B#1'],
    ['6', 'preview 后 status 改变 → commit 拒绝并要求重新 preview', '§C#1'],
    ['7', 'preview 后 research version 改变 → commit 拒绝', '§C#2'],
    ['8', 'preview 后 trade_plan version 改变 → commit 拒绝', '§C#3'],
    ['9', '同一批重复 exchange+code → preview 拒绝', '§D#1'],
    ['10', 'core_validations 为字符串数组 → preview 拒绝', '§E#1'],
    ['11', 'core_validations.status 非法 → preview 拒绝', '§E#2'],
    ['12', 'wall_conditions 为字符串数组 → preview 拒绝', '§E#4'],
    ['13', 'wall_conditions.triggered 非 boolean → preview 拒绝', '§E#5'],
    ['14', 'name mismatch → preview 数据与页面均明显展示警告', '§F#1 / §F#2 / §F#3'],
])
add_p(doc, '另含 §C#4（预览后原本不存在的证券已出现 → 拒绝）、§E#3/§E#6（缺 content → 拒绝）、'
            '§B#2/§B#3（未勾选拒绝 / 勾选后成功且写入台账）、'
            '§G（token 非 base64 / 服务端 cache / TTL 常量 / 版本号静态契约）、'
            '§H（research / plan / execution 版本化与 trades 不变等既有语义不回归）。')

add_p(doc, '')
add_p(doc, '§8.2 既有测试无回退说明', bold=True)
add_p(doc, 'test_v107.py 在 v1.0.7 为 126 断言（本轮在 v1.0.7 原始 ZIP 上实跑核对），'
            '本版为 127 断言，全部通过。其中 5 条被同步改写 —— 这不是新增规则，'
            '而是断言对象本身发生了契约变更；其余 121 条业务语义断言逐字未动：')
add_table(doc, ['断言', 'v1.0.7 期望', 'v1.0.8 期望', '变更原因'], [
    ['§K#7',
     "勾选后 change_confirmed=True（读导入 JSON 的 status_change_confirmed）",
     '拆为 §K#7（粘贴 JSON 自带 confirmed=true → preview 仍 False）'
     ' + §K#7b（该 commit 仍被拒绝）',
     '状态确认不再来自导入 JSON'],
    ['§K#8', 'commit_import_full(token) 成功',
     'commit_import_full(token, [0]) 成功',
     '确认改为 commit 请求单独提交'],
    ['§P#1', "TARGET_SCHEMA_VERSION = '1.0.7'", "TARGET_SCHEMA_VERSION = '1.0.8'",
     '当前版本契约前移'],
    ['§P#2', "server_version = 'Workbench/1.0.7'", "server_version = 'Workbench/1.0.8'",
     '当前版本契约前移'],
    ['§P#6', 'def commit_import_full(token):',
     'def commit_import_full(token, confirmed_status_changes=None):',
     '函数签名随契约变更'],
])
add_p(doc, '净变化：126 → 127（改写 5 条、其中 §K#7 拆为 2 条，其余 121 条原样保留并通过）。'
            '另有一处非断言改动：test_v107.py 增加服务端 preview cache 清理，仅用于测试隔离。')

doc.add_page_break()

# ================ §9 交付包一致性 ================
add_h(doc, '9. 交付包一致性修复（审计问题 9）', 1)
add_p(doc, 'v1.0.7 实际交付 ZIP 无法复现文档宣称的 279/279：ZIP 内缺少 '
            'tests/test_integration_quote.py，导致 test_v102.py 中"集成测试文件存在"断言失败。'
            'v1.0.8 逐项修复：')
add_table(doc, ['项', 'v1.0.7 问题', 'v1.0.8 处理'], [
    ['集成测试文件', 'ZIP 缺 test_integration_quote.py，test_v102.py 解压后失败',
     '纳入 ZIP，解压即可复现 test_v102.py 全绿'],
    ['测试数字', 'PACK_NOTES 声称 279 断言，实际不可复现',
     '实跑后写入真实数字（离线 407 + 集成 14 + 冒烟 14）'],
    ['Manifest 数量', 'Manifest 列 14 项，ZIP 实际 17 项',
     'Manifest 逐项与 ZIP 实际内容一一对应，数量一致'],
    ['Manifest 幽灵条目', '列出了 ZIP 中不存在的 build 脚本',
     '不再列出任何 ZIP 中不存在的脚本；生成脚本不进 Manifest'],
    ['README 地址', 'README 写 http://127.0.0.1:8000，与默认端口 8765 不符',
     '改为真实的 http://127.0.0.1:8765'],
    ['SHA-256', '基于不一致的文件集计算',
     '对最终 ZIP 内全部文件重新计算，写入 Manifest'],
    ['启动脚本', '硬编码本机绝对路径且中文用户名乱码',
     '改为 PATH 查找 python，去除非可移植路径'],
    ['生产数据库', 'schema_version 仍为 1.0.7',
     '迁移至 1.0.8，integrity_check=ok、foreign_key_check 0 违规'],
])

# ================ §10 边界 ================
add_h(doc, '10. 本轮边界（未改变的部分）', 1)
add_bullet(doc, '未新增数据库表 / 列；未改动 create_security_tx / update_research_tx / '
                'update_plan_tx / add_execution_tx 的 SQL 路径。')
add_bullet(doc, '未改变交易规则、状态机、评分口径或研究口径。')
add_bullet(doc, '未新增非导入相关的业务模块或接口。')
add_bullet(doc, '导入失败仍为整批 ROLLBACK；execution 仍 append-only；导入仍不触碰 trades。')
add_bullet(doc, '继续不实现 AI 自由文本解析 / Markdown 解析 / 文件拖拽 / 自动联网补全。')

# ================ §11 风险与开放问题 ================
add_h(doc, '11. 风险与开放问题', 1)
add_table(doc, ['ID', '类别', '说明'], [
    ['R-023', '已修复', 'v1.0.8 修复：preview token 不再可被客户端构造（服务端随机 + cache + TTL + 一次性）'],
    ['R-025', '已修复', 'v1.0.8 修复：preview → commit 之间数据库被并发改动 → 漂移检测拦截并要求重新预览'],
    ['R-026', '新增风险', 'preview cache 为进程内存态；服务重启后未使用的 token 全部失效（用户需重新预览）'],
    ['R-027', '新增风险（本轮验证中发现，未修复）',
     '同一 token 的并发 commit 可重复写入 append-only 的 execution。'
     '_import_cache_take() 只读取 cache 不消费，token 在 commit 成功后（_import_cache_invalidate）'
     '才失效，存在 TOCTOU 窗口；而 spec 要求的 4 类漂移检测（status / research version / '
     'trade_plan version / 原不存在证券已出现）不覆盖"新增 execution"这类 append-only 写入。'
     '复现（本机 3/3 次必然触发，具体次数随线程调度波动）：同一 token 并发 8 次 commit → '
     '返回 200 的有 2~3 次（应为 1 次），execution_reviews 相应多写 1~2 条（应为 +1）。'
     '不变式：**成功次数恒 > 1，append-only 表恒被多写**。'
     '注：新建证券 / 改 status / 改 research 或 trade_plan 的场景下，漂移检测可拦住重复提交'
     '（实测 6 线程并发仅 1 次成功）。'
     '影响面：仅本机单用户工作台的毫秒级重复提交，非网络多用户场景。'
     '建议 v1.0.9 修复：在 cache 中加 per-token"in-flight"标记，take 时原子置位，失败释放、成功失效。'],
    ['R-028', '新增风险（本轮验证中发现，未修复）',
     '并发 commit 可能返回 HTTP 500「database is locked」而非 400。'
     '_mut() 只把 sqlite3.IntegrityError 映射为 400，未捕获 sqlite3.OperationalError。'
     '复现：同 token 并发 8 次 commit → 稳定出现若干次 500（本机 2~5 次，随调度波动）。'
     '不变式：**5xx 必然出现（应为 0）**。'
     '影响面：错误码语义不准确（数据未损坏），仅本机单用户场景可触发。'
     '建议 v1.0.9 修复：在 _mut() 中把 sqlite3.OperationalError 一并映射为 400。'],
    ['R-024', '风险', '大批量 securities（>10k）单事务耗时；当前 SQLite 串行写入'],
    ['R-014', '风险', '单一行情源（Tencent）'],
    ['R-015', '风险', 'HKD 缺汇率需用户录入'],
    ['O-003', '待批准功能', 'trades 冲正机制（A/B/C 三方向待批）'],
    ['O-004', '待批准功能', 'securities.research_pool 字段彻底废弃（SQLite 不支持 DROP COLUMN）'],
])
add_p(doc, 'R-027 / R-028 为本轮交付前独立验证中发现，均**未修改任何代码**：'
            '本轮 spec 明确限定"只修复 v1.0.7 审计发现的 9 项、不得扩展任务"，'
            '故仅如实登记，修复需用户明确批准后另开版本。')

doc.add_page_break()

# ================ §12 第三方审计重点 ================
add_h(doc, '12. 第三方审计重点（v1.0.8）', 1)
add_bullet(doc, 'preview 返回的 token 是否由服务端生成（不应等于客户端粘贴内容的任何编码形式）')
add_bullet(doc, '不调用 preview 直接 commit 是否被拒绝；手工构造 token 是否被拒绝')
add_bullet(doc, '同一 token commit 成功后再次使用是否被拒绝；超时 token 是否被拒绝')
add_bullet(doc, '导入 JSON 内 status_change_confirmed=true 时，preview 返回的 change_confirmed 是否仍为 false')
add_bullet(doc, '未提交 confirmed_status_changes 时 commit 是否 400，且库内 status 不变')
add_bullet(doc, 'preview 后手工改库（status / research version / plan version / 新增证券）再 commit，'
                '是否全部被拒绝并提示重新预览')
add_bullet(doc, '同一批 securities[] 出现重复 exchange+code 时 preview 是否整体拒绝')
add_bullet(doc, 'core_validations / wall_conditions 传字符串数组、非法 status、'
                '非布尔 triggered 时 preview 是否拒绝')
add_bullet(doc, 'name 不一致时 preview 返回体与页面是否都明显提示；commit 后库内 name 是否未被修改')
add_bullet(doc, 'app.js / index.html 是否已不存在「载入示例」按钮')
add_bullet(doc, 'samples/json/ 样例是否为明显虚构标的且带"请勿写入正式库"提示')
add_bullet(doc, 'ZIP 内是否包含 tests/test_integration_quote.py；解压后 test_v102.py 是否全绿')
add_bullet(doc, 'Manifest 列出的文件数量与路径是否与 ZIP 实际内容完全一致（无幽灵条目）')
add_bullet(doc, 'README 默认地址是否为 http://127.0.0.1:8765（与 server.py 默认端口一致）')
add_bullet(doc, 'test_v102/v103/v105/v106/v107/v108 是否全部通过（407 断言，无回退）')

add_p(doc, '')
add_p(doc, '以下两项为本轮已知缺陷（R-027 / R-028），审计方可独立复现，'
            '本版未修复且已在 §11 如实登记：', bold=True)
add_bullet(doc, '同一 preview token 并发 commit 8 次：预期"1 次成功 + 其余 400 + 0 个 5xx"；'
                '实测"成功 2~3 次（append-only 表被多写）+ 若干 400 + 2~5 个 500"。'
                '次数随线程调度波动，但**不变式必然被破坏**：成功次数恒 > 1，5xx 恒 > 0')
add_bullet(doc, '复现脚本（工作区侧，不随交付）：'
                '`skills/ah-workbench-release-verify/scripts/concurrency_probe.py`，'
                '3/3 次试验稳定判定失败')
add_bullet(doc, '上述两项的触发前提均为"毫秒级重复提交"；'
                '顺序重复提交（先成功后失败）行为正确，返回 400 且不重复写入')

add_p(doc, '')
add_p(doc, '第三方审计签字：____________________  日期：__________', bold=True)
add_p(doc, '用户确认签字：  ____________________  日期：__________', bold=True)
add_p(doc, '开发责任人：    ____________________  日期：__________', bold=True)

doc.save(OUT_PATH)
print('v1.0.8 docx 生成：%s' % OUT_PATH)

with open(OUT_PATH, 'rb') as f:
    h = hashlib.sha256()
    for chunk in iter(lambda: f.read(8192), b''):
        h.update(chunk)
print('DOCX_SHA256=%s' % h.hexdigest())
